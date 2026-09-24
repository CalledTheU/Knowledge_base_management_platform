# Author: WangLei
# Email: WangLei1578@outlook.com
# Date: 2026-09-23

import hashlib
import hmac
import json
import logging
import logging.handlers
import os
import re
import secrets
import shutil
import sqlite3
import subprocess
import tempfile
import time
import uuid
import asyncio
from difflib import SequenceMatcher
from contextlib import contextmanager
from pathlib import Path

from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parent
load_dotenv(ROOT / ".env")

from fastapi import Depends, FastAPI, Header, HTTPException, Request
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

from api.documents_router import register_router as register_documents_router
from api.organization_router import register_router as register_organization_router
from schema.document_schema import ACLUpdate
from services.embedding_service import encode_documents, encode_query, cosine, sparse_dot
from services import rag_service

DATA = ROOT / "data"
UPLOADS = DATA / "uploads"
LOGS = DATA / "logs"
DB_PATH = Path(os.getenv("KB_DB_PATH", DATA / "knowledge.db"))
UPLOADS.mkdir(parents=True, exist_ok=True)
LOGS.mkdir(parents=True, exist_ok=True)
logger = logging.getLogger("knowledge_base")
logger.setLevel(logging.INFO)
if not logger.handlers:
    file_handler = logging.handlers.RotatingFileHandler(
        LOGS / "app.log", maxBytes=5 * 1024 * 1024, backupCount=5, encoding="utf-8")
    file_handler.setFormatter(logging.Formatter(
        "%(asctime)s %(levelname)s request_id=%(request_id)s %(message)s"))
    logger.addHandler(file_handler)
app = FastAPI(title="知识库管理平台", version="0.1.0")
PERMISSIONS = {"dashboard.view", "knowledge.manage", "curation.manage", "organization.manage", "chat.access"}
BUILTIN_PERMISSIONS = {
    "user": ["chat.access"],
    "knowledge_admin": ["dashboard.view", "knowledge.manage", "curation.manage", "chat.access"],
    "system_admin": sorted(PERMISSIONS),
    "management": ["dashboard.view", "chat.access"],
}
_chat_streams = {}


def chat_stage(request_id, message):
    logger.info("CHAT_STAGE stage=%s", message, extra={"request_id": request_id})
    queue = _chat_streams.get(request_id)
    if queue:
        queue.put_nowait({"type": "stage", "stage": message})


@contextmanager
def db():
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys=ON")
    try:
        yield con
        con.commit()
    finally:
        con.close()


def password_hash(password, salt=None):
    salt = salt or secrets.token_hex(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), bytes.fromhex(salt), 240000).hex()
    return salt, digest


def init_db():
    with db() as con:
        con.executescript("""
        CREATE TABLE IF NOT EXISTS departments(id TEXT PRIMARY KEY, name TEXT NOT NULL, parent_id TEXT);
        CREATE TABLE IF NOT EXISTS roles(id TEXT PRIMARY KEY, name TEXT UNIQUE NOT NULL, built_in INTEGER NOT NULL DEFAULT 0,
            permissions TEXT NOT NULL DEFAULT '[]');
        CREATE TABLE IF NOT EXISTS users(id TEXT PRIMARY KEY, username TEXT UNIQUE NOT NULL, display_name TEXT NOT NULL,
            password_salt TEXT NOT NULL, password_hash TEXT NOT NULL, department_id TEXT, roles TEXT NOT NULL DEFAULT '[]',
            active INTEGER NOT NULL DEFAULT 1);
        CREATE TABLE IF NOT EXISTS sessions(token TEXT PRIMARY KEY, user_id TEXT NOT NULL REFERENCES users(id), expires REAL NOT NULL);
        CREATE TABLE IF NOT EXISTS documents(id TEXT PRIMARY KEY, title TEXT NOT NULL, filename TEXT NOT NULL,
            category TEXT NOT NULL DEFAULT '', content TEXT NOT NULL, enabled INTEGER NOT NULL DEFAULT 1,
            created REAL NOT NULL, source_object TEXT);
        CREATE TABLE IF NOT EXISTS acl(document_id TEXT NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
            kind TEXT NOT NULL, subject TEXT NOT NULL, PRIMARY KEY(document_id,kind,subject));
        CREATE TABLE IF NOT EXISTS chunks(id TEXT PRIMARY KEY, document_id TEXT NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
            ordinal INTEGER NOT NULL, content TEXT NOT NULL, dense_vector TEXT, sparse_vector TEXT,
            vector_indexed INTEGER NOT NULL DEFAULT 0);
        CREATE TABLE IF NOT EXISTS chats(id TEXT PRIMARY KEY, session_id TEXT NOT NULL, user_id TEXT NOT NULL,
            question TEXT NOT NULL, answer TEXT NOT NULL, retrieved TEXT NOT NULL, allowed TEXT NOT NULL,
            denied TEXT NOT NULL, tokens INTEGER NOT NULL, latency_ms INTEGER NOT NULL, created REAL NOT NULL);
        CREATE TABLE IF NOT EXISTS gaps(id TEXT PRIMARY KEY, question TEXT NOT NULL UNIQUE, user_id TEXT NOT NULL,
            department_id TEXT, count INTEGER NOT NULL DEFAULT 1, created REAL NOT NULL, status TEXT NOT NULL DEFAULT 'open');
        CREATE TABLE IF NOT EXISTS faqs(id TEXT PRIMARY KEY, question TEXT NOT NULL UNIQUE, answer TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'candidate', count INTEGER NOT NULL DEFAULT 1, created REAL NOT NULL,
            cache_enabled INTEGER NOT NULL DEFAULT 0);
        CREATE TABLE IF NOT EXISTS faq_cache(question TEXT PRIMARY KEY, faq_id TEXT NOT NULL REFERENCES faqs(id));
        CREATE TABLE IF NOT EXISTS knowledge_tasks(id TEXT PRIMARY KEY, gap_id TEXT NOT NULL UNIQUE REFERENCES gaps(id) ON DELETE CASCADE,
            title TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'open', created REAL NOT NULL);
        """)
        faq_columns = {r["name"] for r in con.execute("PRAGMA table_info(faqs)")}
        if "cache_enabled" not in faq_columns:
            con.execute("ALTER TABLE faqs ADD COLUMN cache_enabled INTEGER NOT NULL DEFAULT 0")
            con.execute("UPDATE faqs SET cache_enabled=1 WHERE id IN (SELECT faq_id FROM faq_cache) AND status='published'")
        has_permissions = "permissions" in {r["name"] for r in con.execute("PRAGMA table_info(roles)")}
        if not has_permissions:
            con.execute("ALTER TABLE roles ADD COLUMN permissions TEXT NOT NULL DEFAULT '[]'")
        chunk_columns = {r["name"] for r in con.execute("PRAGMA table_info(chunks)")}
        document_columns = {r["name"] for r in con.execute("PRAGMA table_info(documents)")}
        if "dense_vector" not in chunk_columns:
            con.execute("ALTER TABLE chunks ADD COLUMN dense_vector TEXT")
        if "sparse_vector" not in chunk_columns:
            con.execute("ALTER TABLE chunks ADD COLUMN sparse_vector TEXT")
        if "vector_indexed" not in chunk_columns:
            con.execute("ALTER TABLE chunks ADD COLUMN vector_indexed INTEGER NOT NULL DEFAULT 0")
        if "source_object" not in document_columns:
            con.execute("ALTER TABLE documents ADD COLUMN source_object TEXT")
        for role_id, label in [("user", "普通用户"), ("knowledge_admin", "知识管理员"),
                               ("system_admin", "系统管理员"), ("management", "管理层")]:
            permissions = json.dumps(BUILTIN_PERMISSIONS[role_id])
            con.execute("INSERT OR IGNORE INTO roles(id,name,built_in,permissions) VALUES(?,?,1,?)",
                        (role_id, label, permissions))
        built_in_roles = con.execute("SELECT id,permissions FROM roles WHERE id IN ('user','knowledge_admin','system_admin','management')").fetchall()
        if built_in_roles and all(not json.loads(role["permissions"]) for role in built_in_roles):
            for role_id, permissions in BUILTIN_PERMISSIONS.items():
                con.execute("UPDATE roles SET permissions=? WHERE id=?",
                            (json.dumps(permissions), role_id))
        if not con.execute("SELECT 1 FROM departments LIMIT 1").fetchone():
            con.executemany("INSERT INTO departments VALUES(?,?,?)", [
                ("dept-finance", "财务部", None), ("dept-hr", "人力资源部", None),
                ("dept-sales", "业务部", None)])
        if not con.execute("SELECT 1 FROM users LIMIT 1").fetchone():
            for username, name, dept, roles, password in [
                ("admin", "系统管理员", "dept-hr", ["system_admin", "knowledge_admin"], "Admin123!"),
                ("finance", "财务用户", "dept-finance", ["user"], "Finance123!"),
                ("staff", "普通用户", "dept-sales", ["user"], "Staff123!")]:
                salt, digest = password_hash(password)
                con.execute("INSERT INTO users VALUES(?,?,?,?,?,?,?,1)",
                            (str(uuid.uuid4()), username, name, salt, digest, dept, json.dumps(roles)))
        if not con.execute("SELECT 1 FROM documents LIMIT 1").fetchone():
            samples = [
                ("差旅报销标准", "差旅报销标准.md", "财务制度", "员工出差前应提交出差申请。交通费用按实际合规票据报销，住宿标准按职级执行。出差结束后十个工作日内提交报销单。", [("global", "*")]),
                ("高管薪酬与股权激励细则", "高管薪酬制度.md", "人事制度", "高管薪酬资料仅限人力资源部和管理层角色查阅。股权激励归属安排由薪酬委员会审核。", [("department", "dept-hr"), ("role", "management")]),
                ("销售退换货处理规范", "退换货.md", "客服规范", "商品破损时请在签收后48小时内提交订单号、外包装和商品照片，客服核实后办理退款或补发。", [("role", "user")]),
            ]
            for title, filename, category, content, permissions in samples:
                doc_id = str(uuid.uuid4())
                con.execute("INSERT INTO documents(id,title,filename,category,content,enabled,created) VALUES(?,?,?,?,?,1,?)",
                            (doc_id, title, filename, category, content, time.time()))
                con.executemany("INSERT INTO acl VALUES(?,?,?)", [(doc_id, *p) for p in permissions])
                for i, chunk in enumerate(split_text(content)):
                    vector = encode_documents([chunk])[0]
                    con.execute("INSERT INTO chunks(id,document_id,ordinal,content,dense_vector,sparse_vector) VALUES(?,?,?,?,?,?)",
                                (str(uuid.uuid4()), doc_id, i, chunk,
                                 json.dumps(vector["dense"]), json.dumps(vector["sparse"])))
        for row in con.execute("SELECT id,content FROM chunks WHERE dense_vector IS NULL").fetchall():
            vector = encode_documents([row["content"]])[0]
            con.execute("UPDATE chunks SET dense_vector=?,sparse_vector=? WHERE id=?",
                        (json.dumps(vector["dense"]), json.dumps(vector["sparse"]), row["id"]))


def split_text(text, size=700):
    from langchain_text_splitters import RecursiveCharacterTextSplitter

    sections, hierarchy, body = [], [""] * 6, []
    in_fence = False

    def flush():
        content = "\n".join([*(item for item in hierarchy if item), *body]).strip()
        if content:
            sections.append(content)
        body.clear()

    for line in text.splitlines():
        if line.lstrip().startswith((chr(96) * 3, "~~~")):
            in_fence = not in_fence
        heading = None if in_fence else re.match(r"^\s*(#{1,6})\s+.+", line)
        if heading:
            if body:
                flush()
            level = len(heading.group(1))
            hierarchy[level - 1] = line.strip()
            hierarchy[level:] = [""] * (6 - level)
        else:
            body.append(line)
    flush()
    splitter = RecursiveCharacterTextSplitter(
        separators=["\n\n", "\n", "。", "！", "？", ".", ",", " ", ""],
        chunk_size=size, chunk_overlap=min(80, size // 8), keep_separator=True)
    return [chunk for section in sections for chunk in splitter.split_text(section)] or (
        [text.strip()] if text.strip() else [])


def repair_corrupt_documents():
    """Rebuild documents imported with broken PDF text extraction."""
    with db() as con:
        documents = con.execute("SELECT id,filename FROM documents WHERE content LIKE '%�%'").fetchall()
    for document in documents:
        source = next((path for path in UPLOADS.glob(f"{document['id']}.*") if path.is_file()), None)
        if not source:
            continue
        try:
            text = validate_extracted_text(extract_file(document["filename"], source.read_bytes()), document["filename"])
            chunks = split_text(text)
            vectors = encode_documents(chunks)
            with db() as con:
                con.execute("UPDATE documents SET content=? WHERE id=?", (text, document["id"]))
                con.execute("DELETE FROM chunks WHERE document_id=?", (document["id"],))
                con.executemany("INSERT INTO chunks(id,document_id,ordinal,content,dense_vector,sparse_vector) VALUES(?,?,?,?,?,?)", [
                    (str(uuid.uuid4()), document["id"], index, chunk,
                     json.dumps(vector["dense"]), json.dumps(vector["sparse"]))
                    for index, (chunk, vector) in enumerate(zip(chunks, vectors))])
            logger.info("Repaired corrupt document id=%s chunks=%s", document["id"], len(chunks),
                        extra={"request_id": "startup"})
        except Exception:
            logger.exception("Failed to repair corrupt document id=%s", document["id"],
                             extra={"request_id": "startup"})


def sync_pending_chunks():
    with db() as con:
        rows = con.execute(
            "SELECT id,document_id,ordinal,dense_vector,sparse_vector FROM chunks WHERE vector_indexed=0"
        ).fetchall()
    if not rows:
        return
    records = [{
        "chunk_id": row["id"], "document_id": row["document_id"],
        "dense_vector": json.loads(row["dense_vector"]),
        "sparse_vector": json.loads(row["sparse_vector"]) if row["sparse_vector"] else {},
    } for row in rows]
    rag_service.upsert_vectors(records)
    with db() as con:
        con.executemany(
            "UPDATE chunks SET vector_indexed=1 WHERE id=? AND dense_vector=? AND coalesce(sparse_vector,'')=coalesce(?,'')",
            [(row["id"], row["dense_vector"], row["sparse_vector"]) for row in rows])


def row_user(row):
    return {"id": row["id"], "username": row["username"], "display_name": row["display_name"],
            "department_id": row["department_id"], "roles": json.loads(row["roles"]), "active": bool(row["active"])}


def current_user(authorization: str | None = Header(default=None)):
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(401, "请先登录")
    token = authorization[7:]
    with db() as con:
        row = con.execute("SELECT u.* FROM sessions s JOIN users u ON u.id=s.user_id WHERE s.token=? AND s.expires>? AND u.active=1",
                          (token, time.time())).fetchone()
        permissions = set()
        if row:
            for role in con.execute("SELECT permissions FROM roles WHERE id IN (SELECT value FROM json_each(?))",
                                    (row["roles"],)):
                permissions.update(json.loads(role["permissions"]))
    if not row:
        raise HTTPException(401, "登录已失效")
    result = row_user(row)
    result["permissions"] = sorted(permissions)
    return result


def require_admin(user=Depends(current_user)):
    if "organization.manage" not in user["permissions"]:
        raise HTTPException(403, "缺少组织管理权限")
    return user


def require_permission(permission):
    def check(user=Depends(current_user)):
        if permission not in PERMISSIONS or permission not in user["permissions"]:
            raise HTTPException(403, "缺少功能权限")
        return user
    return check


def require_any_permission(*permissions):
    def check(user=Depends(current_user)):
        if not set(permissions).intersection(user["permissions"]):
            raise HTTPException(403, "缺少功能权限")
        return user
    return check


def allowed_for(user, acl):
    # Empty ACL means private; any matching ACL entity grants access.
    return any((kind == "global" and subject == "*") or
               (kind == "department" and subject == user["department_id"]) or
               (kind == "user" and subject == user["id"]) or
               (kind == "role" and subject in user["roles"]) for kind, subject in acl)


class Login(BaseModel):
    username: str
    password: str


class Ask(BaseModel):
    question: str = Field(min_length=1, max_length=4000)
    session_id: str | None = None


class FAQUpdate(BaseModel):
    question: str
    answer: str = ""
    status: str = "candidate"


@app.on_event("startup")
def startup():
    init_db()
    repair_corrupt_documents()


@app.middleware("http")
async def log_unhandled_errors(request: Request, call_next):
    request_id = request.headers.get("x-request-id") or str(uuid.uuid4())
    try:
        response = await call_next(request)
    except Exception:
        logger.exception("Unhandled request error method=%s path=%s", request.method, request.url.path,
                         extra={"request_id": request_id})
        from fastapi.responses import JSONResponse
        response = JSONResponse(status_code=500, content={
            "detail": f"请求失败，错误编号：{request_id}", "request_id": request_id})
    response.headers["X-Request-ID"] = request_id
    return response


@app.get("/api/health")
def health():
    return {"status": "ok"}


@app.post("/api/login")
def login(body: Login):
    with db() as con:
        row = con.execute("SELECT * FROM users WHERE username=? AND active=1", (body.username,)).fetchone()
        if not row:
            raise HTTPException(401, "用户名或密码错误")
        _, digest = password_hash(body.password, row["password_salt"])
        if not hmac.compare_digest(digest, row["password_hash"]):
            raise HTTPException(401, "用户名或密码错误")
        token = secrets.token_urlsafe(32)
        con.execute("INSERT INTO sessions VALUES(?,?,?)", (token, row["id"], time.time() + 86400))
    return {"token": token, "user": row_user(row)}


@app.post("/api/logout")
def logout(user=Depends(current_user), authorization: str = Header()):
    with db() as con:
        con.execute("DELETE FROM sessions WHERE token=?", (authorization[7:],))
    return {"ok": True}


@app.get("/api/me")
def me(user=Depends(current_user)):
    return user


def extract_file(filename, raw):
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md"}:
        return raw.decode("utf-8-sig", errors="replace")
    if suffix == ".pdf":
        command = os.getenv("MINERU_COMMAND", "mineru")
        mineru = command if Path(command).is_file() else shutil.which(command)
        if mineru:
            with tempfile.TemporaryDirectory(prefix="kb-mineru-") as temp_dir:
                source = Path(temp_dir) / Path(filename).name
                output = Path(temp_dir) / "output"
                source.write_bytes(raw)
                try:
                    result = subprocess.run(
                        [mineru, "-p", str(source), "-o", str(output), "--backend", "pipeline"],
                        capture_output=True, text=True, encoding="utf-8", errors="replace",
                        timeout=int(os.getenv("KB_MINERU_TIMEOUT", "900")), check=False)
                except subprocess.TimeoutExpired:
                    logger.warning("MinerU timed out; using PDF text fallback",
                                   extra={"request_id": "mineru"})
                else:
                    markdown = output / source.stem / "auto" / f"{source.stem}.md"
                    if result.returncode == 0 and markdown.is_file():
                        return markdown.read_text(encoding="utf-8", errors="replace")
                    logger.warning("MinerU failed exit_code=%s; using PDF text fallback",
                                   result.returncode, extra={"request_id": "mineru"})
        else:
            logger.info("MinerU CLI unavailable; using PDF text fallback",
                        extra={"request_id": "mineru"})
        from pypdf import PdfReader
        import io
        pages = PdfReader(io.BytesIO(raw)).pages
        text = "\n".join(page.extract_text() or "" for page in pages)
        if text.count("�") > max(3, len(text) // 200):
            text = "\n".join(page.extract_text(extraction_mode="layout") or "" for page in pages)
        return text
    if suffix == ".docx":
        from docx import Document
        import io
        return "\n".join(p.text for p in Document(io.BytesIO(raw)).paragraphs)
    raise HTTPException(415, "仅支持 PDF、Markdown、Word、TXT")


def validate_extracted_text(text, filename):
    text = re.sub(r"\x00", "", text).strip()
    if not text:
        raise HTTPException(400, f"文件未提取到文本: {filename}")
    if text.count("�") > max(3, len(text) // 200):
        raise HTTPException(422, f"文件文本编码异常，无法可靠解析: {filename}")
    return text


def words(text):
    result = set()
    for token in re.findall(r"[\u4e00-\u9fff]+|[a-zA-Z0-9_]+", text.lower()):
        if re.fullmatch(r"[\u4e00-\u9fff]+", token):
            result.update(token[i:i + 2] for i in range(len(token) - 1))
            if len(token) == 1:
                result.add(token)
        elif len(token) > 1:
            result.add(token)
    return result


def question_similarity(left, right):
    """Small dependency-free similarity score for FAQ clustering and cache lookup."""
    left = re.sub(r"\W", "", left.lower())
    right = re.sub(r"\W", "", right.lower())
    if not left or not right:
        return 0.0
    if left == right:
        return 1.0
    left_words, right_words = words(left), words(right)
    word_score = len(left_words & right_words) / max(1, len(left_words | right_words))
    left_bigrams = {left[i:i + 2] for i in range(len(left) - 1)}
    right_bigrams = {right[i:i + 2] for i in range(len(right) - 1)}
    char_score = len(left_bigrams & right_bigrams) / max(1, len(left_bigrams | right_bigrams))
    return max(word_score, char_score, SequenceMatcher(None, left, right).ratio())


def build_grounded_answer(question, ranked):
    """Compress retrieved chunks into relevant evidence instead of dumping whole chunks."""
    qwords = words(question)
    candidates = []
    seen = set()
    for chunk_score, _, _, chunk in ranked[:8]:
        for sentence in re.split(r"(?<=[。！？.!?；;])\s*|\n+", chunk["content"]):
            sentence = re.sub(r"\s+", " ", sentence).strip(" -·")
            if len(sentence) < 12 or sentence in seen:
                continue
            seen.add(sentence)
            overlap = len(qwords & words(sentence))
            if overlap or chunk_score >= 0.4:
                candidates.append((chunk_score + min(overlap, 8) * 0.03, sentence))
    candidates.sort(reverse=True, key=lambda item: item[0])
    selected = []
    total = 0
    for _, sentence in candidates:
        if total + len(sentence) > 1400:
            break
        selected.append(sentence)
        total += len(sentence) + 2
        if len(selected) >= 5:
            break
    return "\n\n".join(selected)


@app.post("/api/chat")
def ask(body: Ask, user=Depends(require_permission("chat.access")), request_id: str | None = Header(default=None, alias="X-Request-ID")):
    started = time.perf_counter()
    request_id = request_id or str(uuid.uuid4())
    question = body.question.strip()
    if not question:
        raise HTTPException(400, "Question cannot be empty")
    external = rag_service.external_enabled()
    try:
        chat_stage(request_id, "正在准备知识库检索")
        with db() as con:
            cached_rows = con.execute("SELECT f.* FROM faq_cache c JOIN faqs f ON f.id=c.faq_id WHERE f.status='published' AND f.cache_enabled=1").fetchall()
            cached = max(cached_rows, key=lambda row: question_similarity(question, row["question"]), default=None)
            if cached and question_similarity(question, cached["question"]) < 0.55:
                cached = None
            rows = [] if external else con.execute(
                "SELECT c.id chunk_id,c.content,c.dense_vector,c.sparse_vector,d.id document_id,d.title,d.filename,a.kind,a.subject "
                "FROM chunks c JOIN documents d ON d.id=c.document_id LEFT JOIN acl a ON a.document_id=d.id "
                "WHERE d.enabled=1 ORDER BY d.created DESC").fetchall()
        candidate_ids = []
        if external:
            chat_stage(request_id, "正在同步待索引知识切片")
            sync_pending_chunks()
            chat_stage(request_id, "正在进行问题向量化")
            direct_ids = rag_service.hybrid_search(encode_query(question))
            chat_stage(request_id, "正在进行 Milvus 混合检索")
            try:
                chat_stage(request_id, "正在生成 HyDE 假设文档")
                hyde = rag_service.generate_hypothetical_document(question)
            except Exception:
                hyde = ""
                logger.exception("HyDE generation failed question_length=%s", len(question),
                                 extra={"request_id": str(uuid.uuid4())})
            hyde_ids = rag_service.hybrid_search(encode_query(question + "\n" + hyde)) if hyde else []
            chat_stage(request_id, "正在进行 RRF 检索结果融合")
            candidate_ids = rag_service.reciprocal_rank_fusion(direct_ids, hyde_ids)
            if candidate_ids:
                placeholders = ",".join("?" for _ in candidate_ids)
                with db() as con:
                    rows = con.execute(
                        "SELECT c.id chunk_id,c.content,c.dense_vector,c.sparse_vector,d.id document_id,d.title,d.filename,"
                        "a.kind,a.subject FROM chunks c JOIN documents d ON d.id=c.document_id "
                        "LEFT JOIN acl a ON a.document_id=d.id "
                        f"WHERE d.enabled=1 AND c.id IN ({placeholders})",
                        candidate_ids).fetchall()
    except Exception:
        request_id = str(uuid.uuid4())
        logger.exception("Chat retrieval failed question_length=%s", len(question), extra={"request_id": request_id})
        raise HTTPException(500, f"问答检索失败，错误编号：{request_id}") from None
    by_doc = {}
    chat_stage(request_id, "正在加载候选文档并进行权限过滤")
    for row in rows:
        item = by_doc.setdefault(row["document_id"], {"title": row["title"], "filename": row["filename"], "acl": [], "chunks": []})
        if row["kind"]:
            item["acl"].append((row["kind"], row["subject"]))
        if not any(chunk["id"] == row["chunk_id"] for chunk in item["chunks"]):
            item["chunks"].append({"id": row["chunk_id"], "content": row["content"],
                                   "vector": json.loads(row["dense_vector"]) if row["dense_vector"] else None,
                                   "sparse": json.loads(row["sparse_vector"]) if row["sparse_vector"] else None})
    ranked = []
    denied = []
    reranked = []
    if external:
        rank_by_id = {chunk_id: index for index, chunk_id in enumerate(candidate_ids)}
        for doc_id, doc in by_doc.items():
            can_read = allowed_for(user, doc["acl"])
            for chunk in doc["chunks"]:
                score = 1 / (60 + rank_by_id.get(chunk["id"], len(candidate_ids)))
                target = ranked if can_read else denied
                target.append((score, doc_id, doc, chunk))
        ranked.sort(reverse=True, key=lambda item: item[0])
        denied.sort(reverse=True, key=lambda item: item[0])
        if denied:
            cached = None
        try:
            chat_stage(request_id, "正在使用 BGE Reranker 重排结果")
            reranked = rag_service.rerank(question, [{
                "chunk_id": chunk["id"], "document_id": doc_id,
                "title": doc["title"], "filename": doc["filename"], "content": chunk["content"],
            } for _, doc_id, doc, chunk in ranked])
            reranked = [item for item in reranked
                        if item["score"] >= float(os.getenv("KB_RERANKER_MIN_SCORE", "0.2"))][:5]
        except Exception:
            request_id = str(uuid.uuid4())
            logger.exception("Chat reranking failed question_length=%s", len(question),
                             extra={"request_id": request_id})
            raise HTTPException(500, f"问答重排失败，错误编号：{request_id}") from None
    else:
        qwords = words(question)
        try:
            query_vector = encode_query(question)
        except Exception:
            request_id = str(uuid.uuid4())
            logger.exception("Chat embedding failed question_length=%s", len(question), extra={"request_id": request_id})
            raise HTTPException(500, f"问题向量化失败，错误编号：{request_id}") from None
        for doc_id, doc in by_doc.items():
            can_read = allowed_for(user, doc["acl"])
            for chunk in doc["chunks"]:
                vector_score = cosine(query_vector["dense"], chunk["vector"]) if chunk["vector"] else 0
                sparse_score = sparse_dot(query_vector["sparse"], chunk["sparse"])
                keyword_score = len(qwords & words(chunk["content"]))
                score = vector_score * 0.8 + min(sparse_score, 1) * 0.15 + min(keyword_score, 5) * 0.01
                if score >= 0.25 or keyword_score:
                    (ranked if can_read else denied).append((score, doc_id, doc, chunk))
        ranked.sort(reverse=True, key=lambda item: item[0])
        denied.sort(reverse=True, key=lambda item: item[0])
    if denied:
        cached = None
    citations = ([{"document_id": item["document_id"], "title": item["title"], "filename": item["filename"],
                   "chunk_id": item["chunk_id"], "excerpt": item["content"][:700]} for item in reranked]
                 if external else
                 [{"document_id": doc_id, "title": doc["title"], "filename": doc["filename"],
                   "chunk_id": chunk["id"], "excerpt": chunk["content"][:700]}
                  for _, doc_id, doc, chunk in ranked[:5]])
    if cached:
        chat_stage(request_id, "命中 FAQ 缓存，正在返回答案")
        answer = cached["answer"]
    elif citations:
        if external:
            try:
                chat_stage(request_id, "正在根据授权切片生成答案")
                answer = rag_service.generate_answer(question, reranked)
            except Exception:
                request_id = str(uuid.uuid4())
                logger.exception("Chat answer generation failed question_length=%s", len(question),
                                 extra={"request_id": request_id})
                raise HTTPException(500, f"答案生成失败，错误编号：{request_id}") from None
        else:
            answer = build_grounded_answer(question, ranked) or citations[0]["excerpt"]
    else:
        answer = "目前没有检索到可用于回答的知识内容。问题已记录到知识缺口，管理员可据此补充资料。"
    if denied:
        answer += "\n\n部分参考资料因权限受限无法展示。"
    elapsed = int((time.perf_counter() - started) * 1000)
    chat_stage(request_id, "问答处理完成")
    chat_id = str(uuid.uuid4())
    with db() as con:
        con.execute("INSERT INTO chats VALUES(?,?,?,?,?,?,?,?,?,?,?)", (chat_id, body.session_id or str(uuid.uuid4()),
            user["id"], question, answer, json.dumps([x[1] for x in ranked + denied]),
            json.dumps([x[1] for x in ranked]), json.dumps([x[1] for x in denied]),
            max(1, (len(question) + len(answer)) // 4), elapsed, time.time()))
        if not citations:
            con.execute("INSERT INTO gaps(id,question,user_id,department_id,created) VALUES(?,?,?,?,?) ON CONFLICT(question) DO UPDATE SET count=count+1",
                        (str(uuid.uuid4()), question, user["id"], user["department_id"], time.time()))
        if not cached:
            candidates = con.execute("SELECT id,question FROM faqs").fetchall()
            exact = next((row for row in candidates if row["question"] == question), None)
            match = exact or max((row for row in candidates if row["question"] != question),
                                 key=lambda row: question_similarity(question, row["question"]), default=None)
            if match and question_similarity(question, match["question"]) >= 0.55:
                con.execute("UPDATE faqs SET count=count+1 WHERE id=?", (match["id"],))
            else:
                con.execute("INSERT INTO faqs(id,question,answer,status,count,created,cache_enabled) VALUES(?,?,?,?,1,?,0)",
                            (str(uuid.uuid4()), question, answer, "candidate", time.time()))
    return {"id": chat_id, "question": question, "answer": answer, "citations": citations,
            "restricted": bool(denied), "faq_hit": bool(cached), "latency_ms": elapsed}


@app.post("/api/chat/stream")
async def ask_stream(body: Ask, user=Depends(require_permission("chat.access"))):
    request_id = str(uuid.uuid4())
    queue = asyncio.Queue()
    _chat_streams[request_id] = queue

    async def run():
        try:
            result = await asyncio.to_thread(ask, body, user, request_id)
            for offset in range(0, len(result["answer"]), 32):
                await queue.put({"type": "token", "text": result["answer"][offset:offset + 32]})
            await queue.put({"type": "answer", "answer": result})
        except Exception as exc:
            logger.exception("Chat stream failed", extra={"request_id": request_id})
            await queue.put({"type": "error", "message": str(exc)})
        finally:
            await queue.put({"type": "done"})

    asyncio.create_task(run())

    async def stream():
        try:
            while True:
                event = await queue.get()
                event["request_id"] = request_id
                yield f"event: chat\ndata: {json.dumps(event, ensure_ascii=False)}\n\n"
                if event["type"] == "done":
                    break
        finally:
            _chat_streams.pop(request_id, None)

    return StreamingResponse(stream(), media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@app.get("/api/dashboard")
def dashboard(user=Depends(require_permission("dashboard.view"))):
    with db() as con:
        return {
            "questions": con.execute("SELECT count(*) FROM chats").fetchone()[0],
            "uv": con.execute("SELECT count(DISTINCT user_id) FROM chats").fetchone()[0],
            "knowledge": con.execute("SELECT count(*) FROM documents").fetchone()[0],
            "tokens": con.execute("SELECT coalesce(sum(tokens),0) FROM chats").fetchone()[0],
            "avg_latency": con.execute("SELECT coalesce(avg(latency_ms),0) FROM chats").fetchone()[0],
            "faq_hits": con.execute("SELECT count(*) FROM chats WHERE question IN (SELECT question FROM faq_cache)").fetchone()[0],
            "top_questions": [dict(r) for r in con.execute("SELECT question,count(*) count FROM chats GROUP BY question ORDER BY count DESC LIMIT 8")],
            "top_knowledge": [dict(r) for r in con.execute("SELECT d.title,count(*) count FROM chats c,json_each(c.allowed) j JOIN documents d ON d.id=j.value GROUP BY d.id ORDER BY count DESC LIMIT 8")],
        }


@app.get("/api/faqs")
def faqs(user=Depends(require_permission("curation.manage"))):
    with db() as con:
        return [dict(r) for r in con.execute("SELECT * FROM faqs ORDER BY count DESC,created DESC")]


@app.put("/api/faqs/{faq_id}")
def update_faq(faq_id: str, body: FAQUpdate, user=Depends(require_permission("curation.manage"))):
    question, answer = body.question.strip(), body.answer.strip()
    if not question:
        raise HTTPException(400, "FAQ question cannot be empty")
    if body.status == "published" and not answer:
        raise HTTPException(400, "Published FAQ requires an answer")
    if body.status not in {"candidate", "published", "rejected"}:
        raise HTTPException(400, "FAQ状态无效")
    with db() as con:
        cur = con.execute("UPDATE faqs SET question=?,answer=?,status=? WHERE id=?", (question, answer, body.status, faq_id))
        if not cur.rowcount:
            raise HTTPException(404, "FAQ不存在")
        con.execute("DELETE FROM faq_cache WHERE faq_id=?", (faq_id,))
        if body.status == "published":
            con.execute("INSERT OR REPLACE INTO faq_cache VALUES(?,?)", (question, faq_id))
            con.execute("UPDATE faqs SET cache_enabled=1 WHERE id=?", (faq_id,))
        else:
            con.execute("UPDATE faqs SET cache_enabled=0 WHERE id=?", (faq_id,))
    return {"ok": True}


@app.put("/api/faqs/{faq_id}/cache")
def update_faq_cache(faq_id: str, body: dict, user=Depends(require_permission("curation.manage"))):
    enabled = bool(body.get("enabled"))
    with db() as con:
        faq = con.execute("SELECT question,status FROM faqs WHERE id=?", (faq_id,)).fetchone()
        if not faq:
            raise HTTPException(404, "FAQ not found")
        if enabled and faq["status"] != "published":
            raise HTTPException(400, "Only published FAQ can enable cache")
        con.execute("UPDATE faqs SET cache_enabled=? WHERE id=?", (int(enabled), faq_id))
        con.execute("DELETE FROM faq_cache WHERE faq_id=?", (faq_id,))
        if enabled:
            con.execute("INSERT OR REPLACE INTO faq_cache VALUES(?,?)", (faq["question"], faq_id))
    return {"ok": True, "enabled": enabled}


@app.get("/api/gaps")
def gaps(user=Depends(require_permission("curation.manage"))):
    with db() as con:
        return [dict(r) for r in con.execute("SELECT g.*,d.name department,t.id task_id FROM gaps g LEFT JOIN departments d ON d.id=g.department_id LEFT JOIN knowledge_tasks t ON t.gap_id=g.id ORDER BY g.count DESC,g.created DESC")]


@app.put("/api/gaps/{gap_id}")
def update_gap(gap_id: str, body: dict, user=Depends(require_permission("curation.manage"))):
    status = body.get("status", "open")
    if status not in {"open", "in_progress", "resolved"}:
        raise HTTPException(400, "状态无效")
    with db() as con:
        cur = con.execute("UPDATE gaps SET status=? WHERE id=?", (status, gap_id))
        if not cur.rowcount:
            raise HTTPException(404, "Knowledge gap not found")
    return {"ok": True}


@app.post("/api/gaps/{gap_id}/task")
def create_gap_task(gap_id: str, user=Depends(require_permission("curation.manage"))):
    with db() as con:
        gap = con.execute("SELECT question FROM gaps WHERE id=?", (gap_id,)).fetchone()
        if not gap:
            raise HTTPException(404, "Knowledge gap not found")
        task = con.execute("SELECT * FROM knowledge_tasks WHERE gap_id=?", (gap_id,)).fetchone()
        if task:
            return dict(task)
        task_id = str(uuid.uuid4())
        con.execute("INSERT INTO knowledge_tasks VALUES(?,?,?,?,?)",
                    (task_id, gap_id, "补充知识：" + gap["question"], "open", time.time()))
        return dict(con.execute("SELECT * FROM knowledge_tasks WHERE id=?", (task_id,)).fetchone())


@app.get("/")
def index():
    return FileResponse(ROOT / "front" / "index.html")


register_documents_router(app, db=db, require_permission=require_permission,
                         extract_file=extract_file, validate_text=validate_extracted_text,
                         split_text=split_text, uploads=UPLOADS)
register_organization_router(app, db=db, current_user=current_user,
                             require_permission=require_permission, require_any_permission=require_any_permission,
                             row_user=row_user, password_hash=password_hash,
                             permissions=PERMISSIONS)

app.mount("/static", StaticFiles(directory=ROOT / "front"), name="static")


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="127.0.0.1", port=8000, reload=True)
